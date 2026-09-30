"""Build the usable_2 reference cases (review 2026-09-29) from IFC evidence and advisor workbooks.

Every builder records per-element evidence. Disagreements between independent sources
(recorded properties, model geometry, expert workbook rows) stop generation unless an
adjudication with its reason is listed here. ``--check`` regenerates every case and
compares it with the stored JSON; nothing here is shown to the agent except the rule text.

    python reference_v2.py                 # write all usable_2 cases
    python reference_v2.py --only tread_width_e342__mall
    python reference_v2.py --check         # regenerate and compare with stored files
"""

from __future__ import annotations

import hashlib
import json
import multiprocessing
from collections import Counter
from pathlib import Path

import numpy as np

import simple_rules as sr
from ground_truth import GroundTruth, load, resolve_model_path

TIER = "usable_2"
REVIEW_DATE = "2026-09-29"
TOLERANCE_MM = 0.5
LEVEL_MERGE_MM = 20.0  # tread faces within this height belong to one step (float noise)
MIN_TREAD_AREA_MM2 = 2e4
TAPERED_RATIO = 0.95  # tread area / bounding-rectangle area below this is a tapered step
ARRIVAL_EXTRA_MM = 50.0  # a top face this much deeper than the treads is the arrival landing edge
REPORT_LINE = "\nReport the GlobalId of every element (modelled as {entity}) that does NOT satisfy this requirement."
REGULATION_HEADER = (
    "\nRegulatory text (the authoritative clauses - use them to interpret the requirement and "
    "its applicability; you must still work out how to read or measure it from the IFC):\n"
)
STAIR_APPLICABILITY = (
    "BCA Approved Document (Acceptable Solutions) V7.08 - applicability:\n"
    "E.2.1 A staircase (including a flight of 2 steps or more) shall provide a safe and suitable "
    "passage for movement of people.\n"
    "E.2.2 A staircase must have - ... (b) landings to break a fall and provide a place for rest; "
    "(c) sufficient width, tread and riser to avoid injury; ...\n"
    "E.2.4 The requirements in paragraph E.2.2(b) and (c) do not apply to a staircase located in "
    "any of the following rooms or spaces:\n"
    "  (a) any equipment or plant room;\n"
    "  (b) any production area of an industrial building;\n"
    "  (c) any attic that - (ii) does not exceed an area of 10 square metres; and (iii) is in a "
    "residential building;\n"
    "  (d) any house built for the owner's own use."
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_meshes(model, elements) -> dict[int, tuple[np.ndarray, np.ndarray]]:
    """World-coordinate meshes in metres for *elements*, keyed by IFC entity id."""
    import ifcopenshell.geom

    settings = ifcopenshell.geom.settings()
    settings.set("use-world-coords", True)
    meshes = {}
    if not elements:
        return meshes
    iterator = ifcopenshell.geom.iterator(settings, model, multiprocessing.cpu_count(), include=list(elements))
    if iterator.initialize():
        while True:
            shape = iterator.get()
            faces = np.array(shape.geometry.faces).reshape(-1, 3)
            if len(faces):
                meshes[shape.id] = (np.array(shape.geometry.verts).reshape(-1, 3), faces)
            if not iterator.next():
                break
    return meshes


def _case(case_id, rule, ifc_rel, clause, complexity, labels, provenance, source, notes, reference) -> GroundTruth:
    return GroundTruth(
        case_id=case_id, rule=rule, ifc_path=ifc_rel, clause=clause, complexity=complexity,
        tolerance_mm=TOLERANCE_MM, labels=dict(sorted(labels.items())),
        violations=sorted(g for g, label in labels.items() if label == "FAIL"),
        provenance=dict(sorted(provenance.items())), source=source, notes=notes,
        reference={"tier": TIER, "review_date": REVIEW_DATE, **reference},
    )


# --------------------------------------------------------------------------- #
# Stair geometry
# --------------------------------------------------------------------------- #
def _tapered_width_at_500(polygon) -> float:
    """Tread width 500 mm from the narrow end (E.3.4.3), across the tread's long axis."""
    from shapely.geometry import LineString

    corners = np.array(polygon.minimum_rotated_rectangle.exterior.coords)[:4]
    first, second = corners[1] - corners[0], corners[2] - corners[1]
    along, across = (first, second) if np.linalg.norm(first) >= np.linalg.norm(second) else (second, first)
    along, across = along / np.linalg.norm(along), across / np.linalg.norm(across)
    centre = np.array(polygon.centroid.coords[0])
    position = (np.array(polygon.exterior.coords) - centre) @ along

    def cut(offset: float) -> float:
        point = centre + along * offset
        return polygon.intersection(LineString([point - across * 1e5, point + across * 1e5])).length

    low_end, high_end = cut(position.min() + 5), cut(position.max() - 5)
    narrow, direction = (position.min(), 1) if low_end <= high_end else (position.max(), -1)
    return cut(narrow + direction * 500.0)


def flight_steps(vertices_m: np.ndarray, faces: np.ndarray) -> dict:
    """Upward step faces of one flight mesh: level, depth and shape, in millimetres."""
    from shapely.geometry import Polygon
    from shapely.ops import unary_union

    vertices = vertices_m * 1000.0
    triangles = vertices[faces]
    normals = np.cross(triangles[:, 1] - triangles[:, 0], triangles[:, 2] - triangles[:, 0])
    length = np.linalg.norm(normals, axis=1)
    upward = triangles[(normals[:, 2] / np.maximum(length, 1e-12) > 0.99) & (length > 1e-6)]
    heights = upward[:, :, 2].mean(axis=1)
    order = np.argsort(heights)
    clusters: list[list[int]] = []
    for index in order:
        if clusters and heights[index] - heights[clusters[-1][-1]] <= LEVEL_MERGE_MM:
            clusters[-1].append(index)
        else:
            clusters.append([index])
    steps = []
    for members in clusters:
        polygon = unary_union([Polygon(upward[i][:, :2]).buffer(0.5) for i in members]).buffer(-0.5)
        if polygon.area < MIN_TREAD_AREA_MM2 or polygon.geom_type != "Polygon":
            continue
        corners = np.array(polygon.minimum_rotated_rectangle.exterior.coords)[:4]
        sides = sorted([np.linalg.norm(corners[1] - corners[0]), np.linalg.norm(corners[2] - corners[1])])
        tapered = polygon.area / max(sides[0] * sides[1], 1e-9) < TAPERED_RATIO
        depth = _tapered_width_at_500(polygon) if tapered else sides[0]
        steps.append({"z_mm": round(float(heights[members].mean()), 1), "depth_mm": round(float(depth), 1),
                      "length_mm": round(float(sides[1]), 1), "tapered": bool(tapered)})
    arrival = False
    if len(steps) >= 2:
        typical = float(np.median([step["depth_mm"] for step in steps[:-1]]))
        arrival = steps[-1]["depth_mm"] > typical + ARRIVAL_EXTRA_MM and not steps[-1]["tapered"]
    return {"steps": steps, "top_is_arrival": arrival,
            "z_min_mm": round(float(vertices[:, 2].min()), 1), "z_max_mm": round(float(vertices[:, 2].max()), 1)}


def geometric_riser_count(geometry: dict, recorded_riser_mm: float | None) -> tuple[int, int] | None:
    """(low, high) risers of one flight from its step levels, including the arrival riser.

    The top face may be the arrival level itself or one riser below it; the two bounds
    coincide when the top face is identified as the deeper arrival edge.
    """
    levels = [step["z_mm"] for step in geometry["steps"]]
    if len(levels) >= 2:
        riser = float(np.median(np.diff(levels)))
        low = round((levels[-1] - levels[0]) / riser) + 1
        return (low, low) if geometry["top_is_arrival"] else (low, low + 1)
    if len(levels) == 1 and recorded_riser_mm:
        count = 1 if levels[0] - geometry["z_min_mm"] < 1.5 * recorded_riser_mm else 2
        return count, count + 1
    return None


def stair_applicability(model, stairs, model_id: str) -> dict[str, tuple[str | None, str]]:
    """E.2.1/E.2.4 exemption per IfcStair.

    v8 reuses the expert-backed decisions of staircase_width_v8; other models are
    judged from building properties and the spaces that contain each staircase.
    """
    import ifcopenshell.util.element as element_util

    from stair_width_reference import applicability_finding, stair_evidence

    if model_id == "stair_v8":
        full = load("staircase_width_v8")
        return {s.GlobalId: (("NA", "staircase_width_v8: " + (full.provenance[s.GlobalId].get("adjudication") or full.provenance[s.GlobalId]["applicability"])["reason"])
                             if full.labels.get(s.GlobalId) == "NA" else (None, "No exemption in staircase_width_v8."))
                for s in stairs}
    owner_built = any(element_util.get_psets(b).get("SGPset_Building", {}).get("OwnerBuiltOwnerStay") is True
                      for b in model.by_type("IfcBuilding"))
    if owner_built:
        return {s.GlobalId: ("NA", "Building is recorded as built for the owner's own use (E.2.4(d)).") for s in stairs}
    parts = [p for s in stairs for p in element_util.get_decomposition(s)] + list(stairs) + model.by_type("IfcSpace")
    meshes = load_meshes(model, parts)
    spaces = [(space, meshes[space.id()][0][meshes[space.id()][1]]) for space in model.by_type("IfcSpace") if space.id() in meshes]
    result = {}
    for stair in stairs:
        if not any(p.id() in meshes for p in element_util.get_decomposition(stair)) and stair.id() not in meshes:
            result[stair.GlobalId] = (None, "No stair geometry; no containing space found.")
            continue
        evidence = stair_evidence(model, meshes, spaces, stair)
        finding, reason = applicability_finding(evidence)
        if finding == "NA" and "E.2.1" not in reason:
            result[stair.GlobalId] = ("NA", reason)
        else:
            names = ", ".join(space["name"] for space in evidence["spaces"]) or "none"
            result[stair.GlobalId] = (None, f"No E.2.4 exemption found (containing spaces: {names}).")
    return result


def stair_flights(model) -> dict:
    """{stair GlobalId: [flights]}; flights without a parent stair are grouped under None."""
    import ifcopenshell.util.element as element_util

    grouped: dict = {}
    for flight in model.by_type("IfcStairFlight"):
        parent = element_util.get_aggregate(flight)
        grouped.setdefault(parent.GlobalId if parent is not None and parent.is_a("IfcStair") else None, []).append(flight)
    return grouped


# --------------------------------------------------------------------------- #
# E.3.5.2 intermediate landing every 18 risers (per flight)
# --------------------------------------------------------------------------- #
LANDING_MAX_RISERS = 18
LANDING_MODELS = ("stair_v8", "mall", "scdf_5")
LANDING_RULE = (
    "Except for spiral staircases, an intermediate landing must be provided at intervals of not "
    "more than 18 risers: no stair flight may rise more than 18 risers without a landing.\n"
    + REGULATION_HEADER + STAIR_APPLICABILITY + "\n"
    "E.3.5 Landing:\n"
    "E.3.5.1 A landing shall be provided at every floor level and door opening.\n"
    "E.3.5.2 Except for spiral staircases, an intermediate landing shall be provided in between "
    "floor levels at intervals of not more than 18 risers.\n"
    "E.3.5.4 ... A winder does not constitute a landing ... (winders are counted as risers).\n"
    "\nLabel every stair flight. A flight is one run between landings or floor levels; its riser "
    "count includes the final riser up to the landing or floor it arrives at. Label a flight NA "
    "when it belongs to a spiral staircase or to a staircase exempt under E.2.1 or E.2.4. Building use "
    "is not recorded in these models: a space named as a production area is treated as a production area "
    "of an industrial building (E.2.4(b)). A flight "
    "whose riser count cannot be determined cannot be judged: label it ALERT, and do NOT report it "
    "as a violation."
    + REPORT_LINE.format(entity="IfcStairFlight")
)
# Reviewed differences between recorded and geometric riser counts (none required so far).
LANDING_ADJUDICATIONS: dict[str, dict] = {}


def _flight_count_decision(flights, recorded, geometric) -> dict[str, dict]:
    """Choose each flight's riser count from recorded values and geometric (low, high) bounds."""
    values = [recorded[f.GlobalId] for f in flights]
    bounds = [geometric[f.GlobalId] for f in flights]
    decisions = {}
    over = lambda count: count > LANDING_MAX_RISERS  # noqa: E731
    if len(flights) > 1 and None not in values and len(set(values)) == 1 and None not in bounds:
        total = int(values[0])
        if sum(b[0] for b in bounds) <= total <= sum(b[1] for b in bounds) and total > max(b[1] for b in bounds):
            if any(over(b[0]) != over(b[1]) for b in bounds):
                raise ValueError(f"Staircase total {total} on flights {[f.GlobalId for f in flights]}: bounds straddle the limit. Adjudicate.")
            return {f.GlobalId: {
                "count": geometric[f.GlobalId][1], "source": "geometry (upper bound)",
                "note": f"Recorded {total} is the staircase total repeated on each of {len(flights)} flights; the flights' geometric bounds {[list(b) for b in bounds]} sum to a range containing it."}
                for f in flights}
    for flight in flights:
        guid, value, bound = flight.GlobalId, recorded[flight.GlobalId], geometric[flight.GlobalId]
        if guid in LANDING_ADJUDICATIONS:
            decisions[guid] = {"count": LANDING_ADJUDICATIONS[guid]["count"], "source": "adjudication",
                               "note": LANDING_ADJUDICATIONS[guid]["reason"]}
        elif value is None and bound is None:
            decisions[guid] = {"count": None, "source": "none", "note": "No recorded or measurable riser count."}
        elif bound is None:
            decisions[guid] = {"count": int(value), "source": "recorded only", "note": ""}
        elif value is None:
            if over(bound[0]) != over(bound[1]):
                raise ValueError(f"{guid}: no recorded count and geometric bounds {bound} straddle the limit. Adjudicate.")
            decisions[guid] = {"count": bound[0], "source": "geometry only", "note": f"Geometric bounds {bound}."}
        elif bound[0] <= int(value) <= bound[1]:
            decisions[guid] = {"count": int(value), "source": "recorded, within geometric bounds", "note": ""}
        elif over(int(value)) == over(bound[0]) == over(bound[1]):
            decisions[guid] = {"count": int(value), "source": "recorded; geometry gives the same outcome",
                               "note": f"Geometric bounds {bound}."}
        else:
            raise ValueError(f"{guid}: recorded {value} and geometric {bound} risers give different outcomes. Adjudicate.")
    return decisions


def survey_stairs(model, model_id: str) -> dict:
    """Per-flight recorded values, step geometry, riser-count decisions and stair applicability."""
    scale = sr._mm_scale(model)
    grouped = stair_flights(model)
    stairs = [model.by_guid(guid) for guid in grouped if guid is not None]
    applicability = stair_applicability(model, stairs, model_id)
    flights = [f for group in grouped.values() for f in group]
    meshes = load_meshes(model, flights)
    recorded, geometric, evidence, steps = {}, {}, {}, {}
    for flight in flights:
        value = sr._lookup(flight, sr.RISER_COUNT_KEYS)
        riser = sr._lookup(flight, sr.RISER_KEYS, scale)
        geometry = flight_steps(*meshes[flight.id()]) if flight.id() in meshes else {"steps": [], "top_is_arrival": False, "z_min_mm": None, "z_max_mm": None}
        bounds = geometric_riser_count(geometry, riser)
        recorded[flight.GlobalId], geometric[flight.GlobalId], steps[flight.GlobalId] = value, bounds, geometry
        evidence[flight.GlobalId] = {"recorded_riser_count": value, "recorded_riser_height_mm": None if riser is None else round(riser, 1),
                                     "step_levels": len(geometry["steps"]), "top_face_is_arrival": geometry["top_is_arrival"],
                                     "geometric_riser_bounds": None if bounds is None else list(bounds)}
    stairs_out = {}
    for stair_guid, group in grouped.items():
        stair = model.by_guid(stair_guid) if stair_guid else None
        decisions = _flight_count_decision(group, recorded, geometric)
        stairs_out[stair_guid] = {
            "flights": group, "shape": str(stair.PredefinedType) if stair is not None else "no parent IfcStair",
            "applicability": applicability.get(stair_guid, (None, "No parent IfcStair.")), "decisions": decisions,
            "risers": sum(d["count"] for d in decisions.values() if d["count"] is not None)}
    return {"stairs": stairs_out, "evidence": evidence, "steps": steps, "scale": scale}


def stair_exemption(stair: dict) -> str | None:
    """NA reason shared by the stair rules (E.2.4 exemption or E.2.1 single step), else None."""
    exemption, reason = stair["applicability"]
    if exemption == "NA":
        return reason
    if stair["risers"] < 2:
        return "The staircase rises by a single step; fewer than 2 steps is not a staircase (E.2.1)."
    return None


def build_landing(model_id: str) -> GroundTruth:
    import ifcopenshell

    context = sr.MODELS[model_id]
    path = resolve_model_path(context.ifc_rel)
    survey = survey_stairs(ifcopenshell.open(str(path)), model_id)
    labels, provenance = {}, {}
    for stair_guid, stair in survey["stairs"].items():
        exempt = stair_exemption(stair)
        for flight in stair["flights"]:
            guid, decision = flight.GlobalId, stair["decisions"][flight.GlobalId]
            if exempt:
                label, reason = "NA", exempt
            elif stair["shape"] == "SPIRAL_STAIR":
                label, reason = "NA", "Flight of a spiral staircase (E.3.5.2 exception)."
            elif decision["count"] is None:
                label, reason = "ALERT", decision["note"]
            else:
                label = "FAIL" if decision["count"] > LANDING_MAX_RISERS else "PASS"
                reason = f"{decision['count']} risers in the flight ({decision['source']})."
            labels[guid] = label
            provenance[guid] = {**survey["evidence"][guid], "stair": stair_guid, "stair_type": stair["shape"],
                                "flights_in_stair": len(stair["flights"]), "applicability": stair["applicability"][1],
                                "riser_count": decision, "label": label, "reason": reason}
    counts = Counter(labels.values())
    return _case(
        f"landing_interval_e352__{model_id}", LANDING_RULE, context.ifc_rel, "e.3.5.2", 2, labels, provenance,
        "reference_v2.py::build_landing",
        f"E.3.5.2 per flight with spiral and E.2.1/E.2.4 applicability: {dict(sorted(counts.items()))}. "
        f"Supersedes landing_interval_e352__{model_id}_deprecated_v1, which applied staircase totals to flights and ignored the spiral exception.",
        {"version": 2, "supersedes": f"landing_interval_e352__{model_id}_deprecated_v1", "evaluation_scope": "full_population",
         "ifc_sha256": sha256(path), "population": len(labels),
         "status": "IFC evidence: recorded riser counts corroborated by step geometry; applicability from building properties and containing spaces. Not expert certification.",
         "policy": "Riser count per flight includes the arrival riser. A recorded count equal on every flight of a multi-flight stair and equal to the sum of their geometric counts is treated as the staircase total. Spiral staircase flights and E.2.1/E.2.4-exempt staircases are NA.",
         "limitations": "Building use is not recorded in every model; E.2.4 relies on space names/occupancy and building properties. Geometric counts come from tread-face levels."},
    )


# --------------------------------------------------------------------------- #
# E.3.4.2 / E.3.4.3 tread width, mall (measured per tread)
# --------------------------------------------------------------------------- #
TREAD_RULE = (
    "The width of a tread must not be less than the minimum required for this type of building.\n"
    + REGULATION_HEADER + STAIR_APPLICABILITY + "\n"
    "E.3.4 Risers and treads:\n"
    "E.3.4.2 The width of a tread of a staircase (see Figure E.3.4.2(a) for measurement of "
    "\u201ctread\u201d) shall not be less than:\n"
    "  a) 225 mm, if the staircase is in a residential unit within a residential building;\n"
    "  b) 250 mm, if the staircase is in an industrial building; or\n"
    "  c) 275 mm, if the staircase is in any other type of building, including common staircases "
    "in a residential building.\n"
    "E.3.4.3 The width of the tread of any tapered step shall be measured at a distance of 500 mm "
    "from the narrower end.\n"
    "\nThis model is {building}. Label every stair flight: it fails when any of its treads is "
    "narrower than the minimum. Label a flight NA when its staircase is exempt under E.2.1 or E.2.4, "
    "or when the flight has only one riser, so that no tread lies between its risers; a flight with "
    "two or more risers has at least one tread and is assessed. "
    "A flight whose tread width cannot be determined cannot be judged: label it ALERT, and do NOT "
    "report it as a violation."
    + REPORT_LINE.format(entity="IfcStairFlight")
)


def build_tread_mall() -> GroundTruth:
    import ifcopenshell

    context = sr.MODELS["mall"]
    minimum = sr.TREAD_MIN_MM[context.building_type]
    path = resolve_model_path(context.ifc_rel)
    survey = survey_stairs(ifcopenshell.open(str(path)), "mall")
    labels, provenance = {}, {}
    for stair_guid, stair in survey["stairs"].items():
        exempt = stair_exemption(stair)
        for flight in stair["flights"]:
            guid, geometry = flight.GlobalId, survey["steps"][flight.GlobalId]
            treads = geometry["steps"][:-1] if geometry["top_is_arrival"] else geometry["steps"]
            recorded = sr._lookup(flight, sr.TREAD_KEYS, survey["scale"])
            recorded = None if recorded is None else round(recorded, 1)
            measured = min((t["depth_mm"] for t in treads), default=None)
            risers = stair["decisions"][guid]["count"]
            if exempt:
                label, reason = "NA", exempt
            elif risers == 1:
                label, reason = "NA", "Single-riser flight: it has no tread between risers; its only horizontal face is the landing edge."
            elif measured is None:
                label, reason = "ALERT", "No measurable tread."
            else:
                label = "PASS" if measured >= minimum - TOLERANCE_MM else "FAIL"
                reason = f"Narrowest tread {measured} mm against {minimum:g} mm."
                if recorded is not None and (recorded >= minimum - TOLERANCE_MM) != (label == "PASS"):
                    raise ValueError(f"{guid}: recorded {recorded} mm and measured {measured} mm give different outcomes. Adjudicate.")
                if recorded is not None and abs(recorded - measured) > 1.0:
                    reason += f" Recorded value {recorded} mm differs but gives the same outcome."
            labels[guid] = label
            provenance[guid] = {"recorded_tread_mm": recorded, "measured_min_tread_mm": measured, "treads_measured": len(treads),
                                "tapered_treads": sum(t["tapered"] for t in treads), "flight_risers": risers,
                                "stair": stair_guid, "staircase_risers": stair["risers"], "applicability": stair["applicability"][1],
                                "minimum_mm": minimum, "label": label, "reason": reason}
    counts = Counter(labels.values())
    return _case(
        "tread_width_e342__mall", TREAD_RULE.format(building=context.building_phrase), context.ifc_rel, "e.3.4.2", 1,
        labels, provenance, "reference_v2.py::build_tread_mall",
        f"E.3.4.2/E.3.4.3 measured on every tread, with E.2.1/E.2.4 applicability: {dict(sorted(counts.items()))}. Supersedes tread_width_e342__mall_deprecated_v1.",
        {"version": 2, "supersedes": "tread_width_e342__mall_deprecated_v1", "evaluation_scope": "full_population",
         "ifc_sha256": sha256(path), "population": len(labels),
         "status": "Geometric tread depths (rectangular treads: short side; tapered treads: 500 mm from the narrow end) corroborated by recorded tread values; outcome conflicts stop generation. Not expert certification.",
         "policy": f"Minimum {minimum:g} mm for a building that is neither industrial nor residential. A flight fails if any tread is narrower. Single-step staircases (E.2.1) are NA.",
         "limitations": "Tread faces are read from the exported flight meshes; the arrival edge is excluded when it is deeper than the treads."},
    )


# --------------------------------------------------------------------------- #
# Advisor workbook cases (expert rows, re-scoped and cross-checked)
# --------------------------------------------------------------------------- #
BCA_CASES_REL = "bca_tests-harness_improvements/bca_tests-harness_improvements/BCATestCases"


def workbook_rows(workbook: str, sheet: str) -> dict[str, dict]:
    """Authored expert rows of one comparison sheet, keyed by GlobalId."""
    import openpyxl

    from import_bca_gt import GT_COMPARISON_DIR

    rows = list(openpyxl.load_workbook(GT_COMPARISON_DIR / workbook, data_only=True)[sheet].iter_rows(values_only=True))
    header = [str(cell).strip().lower() for cell in rows[0]]
    out = {}
    for row in rows[1:]:
        record = dict(zip(header, row))
        bucket = str(record["gt_bucket"] or "blank").strip().lower()
        if bucket == "blank":
            continue
        if record["guid"] in out:
            raise ValueError(f"{record['guid']}: several authored rows in {sheet}. Review required.")
        out[record["guid"]] = {"clause": record["clause"], "bucket": bucket, "message": str(record["gt_message"] or "")}
    return out


def template_requirement(folder: str, template: str) -> str:
    """Full requirement text from the authority's template (the comparison workbook truncates it)."""
    import openpyxl

    sheet = openpyxl.load_workbook(resolve_model_path(f"{BCA_CASES_REL}/{folder}/Template/{template}"), data_only=True).worksheets[0]
    texts = {str(row[4]).strip() for row in list(sheet.iter_rows(values_only=True))[1:] if row[4]}
    if len(texts) != 1:
        raise ValueError(f"{template}: expected one requirement text, found {len(texts)}.")
    return texts.pop()


def _space_name(space) -> str:
    import ifcopenshell.util.element as element_util

    return str(element_util.get_psets(space).get("SGPset_Space", {}).get("SpaceName") or space.LongName or space.Name or "")


EXIT_FOLDER = "Location of exits & access to exits (EIB related)"
EXIT_MODELS = {"scdf_5": "uat2_scdf_2.2.12_5_v0", "scdf_3": "uat2_scdf_2.2.12_3_v1"}
EXIT_RULE = (
    "Every occupant must have direct access to the required exits without passing through another "
    "room or tenanted space; where that is not possible, the escape route may pass through at most one "
    "other room of the same tenant, through an exit access door fitted with an electromechanical device.\n"
    + REGULATION_HEADER.replace("the IFC):", "the IFC):\nSCDF Fire Code 2023, clause 2.2.12:") + "{requirement}\n"
    "\nExemptions (authority summary): dwelling units in Purpose Group I and II buildings; "
    "non-functional spaces and enclosed spaces (no door or opening); spaces provided with only one exit "
    "towards exit facilities.\n"
    "Authority interpretation: the clause detects 'room-to-room' issues. A room-to-room issue occurs when "
    "more than one exit is provided from a room or space and some of these exits could be locked while the "
    "room or space is still accessible through the remaining door(s). An escape route has no room-to-room "
    "issue when (a) it opens into a common area such as a corridor or lobby; (b) a room or space opens into "
    "a reception area; or (c) it moves from one room to another room with only one egress route. A "
    "reception or tenanted corridor that does not share its space with another room is not counted as a "
    "room. Condition 2 (obstruction of the doorway) is not part of this check.\n"
    "Authority IFC conventions: exit facilities are IfcSpace with SGPset_Space.SpaceName "
    "InternalExitStaircase, ExternalExitStaircase, InternalExitPassageway, ExternalExitPassageway, "
    "InternalExitRamp or ExternalExitRamp; an exit access is an IfcDoor with Pset_DoorCommon.FireExit = "
    "TRUE or an IfcOpeningElement of type OPENING; an exit access door with an electromechanical device has "
    "SGPset_Door.ElectromechanicalLockingDevice = TRUE; egress indicator boxes "
    "(IfcBuildingElementProxy EGRESSINDICATORBOX) mark the egress side of exit access doors. Common areas "
    "are corridors, lobbies and similar shared circulation spaces (the authority's space classification "
    "list is not supplied).\n"
    "\nLabel every space: PASS when every escape route of its occupants complies, FAIL when a route has a "
    "room-to-room issue that the clause does not permit, NA when the space is exempt. Spaces that are "
    "themselves common areas or exit facilities may be labelled NA."
    + REPORT_LINE.format(entity="IfcSpace")
)


def build_exit(model_id: str) -> GroundTruth:
    import ifcopenshell
    import ifcopenshell.util.element as element_util

    sheet = EXIT_MODELS[model_id]
    context = sr.MODELS[model_id]
    path = resolve_model_path(context.ifc_rel)
    model = ifcopenshell.open(str(path))
    experts = workbook_rows("exit_requirements.xlsx", sheet)
    exits: dict[int, list] = {}
    for boundary in model.by_type("IfcRelSpaceBoundary"):
        element, space = boundary.RelatedBuildingElement, boundary.RelatingSpace
        if element is not None and space is not None and space.is_a("IfcSpace") and (element.is_a("IfcDoor") or element.is_a("IfcOpeningElement")):
            exits.setdefault(element.id(), set()).add(space.id())
    labels, provenance = {}, {}
    for guid, expert in sorted(experts.items()):
        space = model.by_guid(guid)
        if expert["clause"] != "fc:2.2.12b" or expert["bucket"] not in ("pass", "fail"):
            raise ValueError(f"{guid}: unmapped exit row {expert}. Review required.")
        exempt = "exempt" in expert["message"].lower()
        label = "FAIL" if expert["bucket"] == "fail" else "NA" if exempt else "PASS"
        doors = []
        for element_id, spaces in exits.items():
            if space.id() in spaces:
                element = model.by_id(element_id)
                psets = element_util.get_psets(element)
                doors.append({"element": element.GlobalId, "class": element.is_a(),
                              "fire_exit": psets.get("Pset_DoorCommon", {}).get("FireExit"),
                              "electromechanical_device": psets.get("SGPset_Door", {}).get("ElectromechanicalLockingDevice"),
                              "to": sorted(_space_name(model.by_id(other)) for other in spaces if other != space.id())})
        if label == "FAIL" and not any(d["electromechanical_device"] is not True and any(
                not any(word in name.lower() for word in ("corridor", "lobby", "driveway", "exit")) for name in d["to"]) for d in doors):
            raise ValueError(f"{guid}: expert FAIL but no room-to-room door without an electromechanical device. Adjudicate.")
        labels[guid] = label
        provenance[guid] = {"space": _space_name(space), "expert": expert, "boundary_doors": sorted(doors, key=lambda d: d["element"]),
                            "label": label, "mapping": "Expert Pass whose explanation states an exemption is NA." if exempt and label == "NA" else "Expert verdict."}
    population = {space.GlobalId for space in model.by_type("IfcSpace")}
    if set(labels) - population:
        raise ValueError("Expert rows reference GlobalIds that are not IfcSpace in the model.")
    unscored = sorted(population - set(labels))
    counts = Counter(labels.values())
    rule = EXIT_RULE.format(requirement=template_requirement(EXIT_FOLDER, sheet + ".xlsx"))
    return _case(
        f"exit_requirements__{model_id}", rule, context.ifc_rel, "fc:2.2.12b", 3, labels, provenance,
        f"reference_v2.py::build_exit (exit_requirements.xlsx::{sheet})",
        f"Expert rows for the {len(labels)} spaces listed in the authority template: {dict(sorted(counts.items()))}; "
        f"{len(unscored)} other IfcSpace are unscored. " + ("Supersedes exit_requirements__scdf_5_deprecated_v1 (truncated rule, no declared subset)." if model_id == "scdf_5" else "Control: no failures."),
        {"version": 2, "supersedes": "exit_requirements__scdf_5_deprecated_v1" if model_id == "scdf_5" else None,
         "evaluation_scope": "annotated_subset", "unscored_ids": unscored, "population": len(population),
         "ifc_sha256": sha256(path),
         "status": "Advisor expert rows; FAIL rows cross-checked against space-boundary doors and electromechanical-device properties. Not independent expert certification.",
         "policy": "Scope is the authority template's spaces; other IfcSpace are unscored. Expert Pass rows whose explanation states an exemption (one exit, non-functional space) are NA. Full clause text from the template.",
         "limitations": "The advisors' own checker agreed with 44 of 51 exit rows; the tenant-internal corridor FAIL (0eIFN9_CjDohnfRZazdLH9, ...LK5) is an interpretation of the authority's reception/transitional-space rule."},
    )


PLANTER_RULE = (
    "A communal planter box is excluded from Gross Floor Area (GFA) only if it meets every requirement "
    "below; otherwise the entire planter box is included in GFA (there is no partial exclusion).\n"
    + REGULATION_HEADER + "URA GFA Handbook - planter boxes (authority summary):\n"
    "Communal planter boxes are excluded from GFA if they comply with the following requirements: "
    "maximum width of 1.0 m; minimum depth of 500 mm (for planting purposes); to be sufficiently "
    "externalised. Planter boxes provided behind full height screening and internalised planter boxes are "
    "included as GFA. A planter box that is not within a common or communal area is included as GFA.\n"
    "Authority definitions: planter box - IfcFurniture USERDEFINED PLANTERBOX; communal use - IfcSpace "
    "USERDEFINED AREA_STRATA whose SGPset_SpaceArea_Strata.AST_AreaType contains 'Common' or 'Communal'; "
    "communal planter box - any part of the planter box overlaps a communal space in plan; width - the "
    "length of the side measured from the edge near the building to the edge away from the building (if "
    "the width varies, no width may exceed 1 m); depth - the overall height of the planter box (if the "
    "depth varies, no depth may be less than 500 mm); externalised - the sides of the planter box facing "
    "out are not obstructed by any building within 20 m. Suggested externalisation check: place 2 layers "
    "of 3 points along the building-edge length of the planter's bounding box, at 1100 mm and 1500 mm "
    "above the planter surface; project 10 rays outwards from each point for 20 m; the planter is "
    "externalised if at least 40% of the rays at each layer are not obstructed by other physical elements.\n"
    "\nLabel every planter box: PASS when it qualifies for exclusion from GFA, FAIL when it must be "
    "included in GFA."
    + REPORT_LINE.format(entity="IfcFurniture")
)


def build_planter() -> GroundTruth:
    import re

    import ifcopenshell
    import ifcopenshell.util.element as element_util
    import shapely
    from shapely.geometry import MultiPoint, Polygon
    from shapely.ops import unary_union

    context = sr.MODELS["planterbox"]
    path = resolve_model_path(context.ifc_rel)
    model = ifcopenshell.open(str(path))
    experts = workbook_rows("planter_box.xlsx", "20240313_URA_Planterbox")
    planters = model.by_type("IfcFurniture")
    strata = [s for s in model.by_type("IfcSpace") if "STRATA" in (s.ObjectType or "").upper()]
    meshes = load_meshes(model, planters + strata)
    if {p.GlobalId for p in planters} != set(experts):
        raise ValueError("Expert rows do not cover exactly the IfcFurniture population.")
    labels, provenance = {}, {}
    for planter in planters:
        expert = experts[planter.GlobalId]
        vertices, faces = meshes[planter.id()]
        triangles = vertices[faces] * 1000
        footprint = unary_union([Polygon(t[:, :2]).buffer(0.5) for t in triangles if Polygon(t[:, :2]).area > 1]).buffer(-0.5)
        outline = unary_union([Polygon(g.exterior) for g in getattr(footprint, "geoms", [footprint])])
        widest = 2 * shapely.maximum_inscribed_circle(outline, tolerance=1.0).length
        communal = sorted({str(element_util.get_psets(s).get("SGPset_SpaceArea_Strata", {}).get("AST_AreaType"))
                           for s in strata if s.id() in meshes
                           and MultiPoint(meshes[s.id()][0][:, :2] * 1000).convex_hull.intersection(outline).area > 1e4})
        message = expert["message"]
        stated = {key: float(value) for key, value in re.findall(r"(depth|width)=(\d+)mm", message)}
        measured_width_ok = widest <= 1000 + TOLERANCE_MM
        if measured_width_ok != (stated["width"] <= 1000):
            raise ValueError(f"{planter.GlobalId}: measured widest strip {widest:.0f} mm and expert width {stated['width']:.0f} mm disagree on the 1 m limit.")
        if ("Communal" in message) != any("communal" in c.lower() or "common" in c.lower() for c in communal):
            raise ValueError(f"{planter.GlobalId}: communal overlap disagrees with the expert row.")
        label = {"pass": "PASS", "fail": "FAIL"}[expert["bucket"]]
        labels[planter.GlobalId] = label
        provenance[planter.GlobalId] = {
            "expert": expert, "measured_widest_strip_mm": round(widest), "mesh_height_mm": round(float(np.ptp(triangles[:, :, 2]))),
            "overlapping_strata": communal, "label": label,
            "check": "Width-limit outcome and communal overlap re-measured and agree with the expert row. The expert depth (overall height) exceeds the mesh height of the planter body; no label depends on which reading is used."}
    if labels.get("1htswfSyH2TvqYjduovEU2") != "PASS" or provenance["1htswfSyH2TvqYjduovEU2"]["mesh_height_mm"] < 500:
        raise ValueError("The only excluded planter must meet the 500 mm depth by both readings.")
    counts = Counter(labels.values())
    return _case(
        "planter_box_gfa", PLANTER_RULE, context.ifc_rel, "-", 3, labels, provenance,
        "reference_v2.py::build_planter (planter_box.xlsx::20240313_URA_Planterbox)",
        f"GFA-exclusion eligibility of all 8 planter boxes: {dict(sorted(counts.items()))}. PASS = excluded from GFA. Supersedes planter_box_gfa_deprecated_v1, which gave no criteria.",
        {"version": 2, "supersedes": "planter_box_gfa_deprecated_v1", "evaluation_scope": "full_population", "population": len(labels),
         "ifc_sha256": sha256(path),
         "status": "Advisor expert rows; width outcome and communal overlap re-measured from the IFC. Externalisation is not independently re-measured.",
         "policy": "The rule states the URA criteria and the authority's IFC definitions; FAIL means the planter must be included in GFA.",
         "limitations": "Eight elements only. The advisors' checker disagreed on 1 of 8 (width axis of the L-shaped excluded planter); the measured widest strip there is 1000 mm."},
    )


HEADROOM_DOOR_RULE = (
    "The headroom of every room, access route and circulation space must be at least 2.0 m; at a doorway "
    "it is measured to the underside of the transom.\n"
    + REGULATION_HEADER + "BCA Approved Document (Acceptable Solutions) V7.08:\n"
    "C.3.2.1 The headroom of every room, access route and circulation space shall not be less than 2.0 metres.\n"
    "Note: 1 The term \u201caccess route\u201d shall include a covered walkway or footway of a building. "
    "2 The headroom is measured from the finished floor level to - a) in the case of a doorway, the "
    "underside of the transom; ... c) in all other cases, the underside of any beam, duct, service pipe, "
    "fixture, fitting or other obstruction or projection.\n"
    "Authority IFC convention (headroom test summary): the underside of the transom is the door's recorded "
    "ClearHeight in its IFC-SG door property set.\n"
    "Benchmark policy: every doorway is assessed; the room exemptions (attics, plant rooms, small stores, "
    "owner-built houses) do not exempt a doorway.\n"
    "\nLabel every door."
    + REPORT_LINE.format(entity="IfcDoor")
)


def build_headroom_doors() -> GroundTruth:
    import ifcopenshell
    import ifcopenshell.util.element as element_util

    context = sr.MODELS["mall"]
    path = resolve_model_path(context.ifc_rel)
    model = ifcopenshell.open(str(path))
    scale = sr._mm_scale(model)
    experts = {g: e for g, e in workbook_rows("headroom.xlsx", "Headroom_beta mall_UAT1_v6.4").items() if e["clause"] == "c.3.2"}
    labels, provenance = {}, {}
    for door in model.by_type("IfcDoor"):
        expert = experts.get(door.GlobalId)
        if expert is None:
            continue
        clear = element_util.get_psets(door).get("SGPset_Door", {}).get("ClearHeight")
        clear_mm = None if clear is None else round(clear * scale, 1)
        label = {"pass": "PASS", "fail": "FAIL"}[expert["bucket"]]
        if clear_mm is None or (clear_mm >= 2000 - TOLERANCE_MM) != (label == "PASS"):
            raise ValueError(f"{door.GlobalId}: recorded ClearHeight {clear_mm} disagrees with the expert {label}.")
        overall = None if door.OverallHeight is None else round(door.OverallHeight * scale, 1)
        labels[door.GlobalId] = label
        provenance[door.GlobalId] = {
            "expert": expert, "clear_height_mm": clear_mm, "overall_height_mm": overall, "label": label,
            "conflict": None if overall is None or (overall >= 2000 - TOLERANCE_MM) == (label == "PASS") else
            "OverallHeight gives the opposite outcome; the authority convention (ClearHeight) governs."}
    if any(not model.by_guid(guid).is_a("IfcDoor") and not model.by_guid(guid).is_a("IfcStair") for guid in set(experts) - set(labels)):
        raise ValueError("Expert c.3.2 rows reference elements other than doors and stairs.")
    population = {door.GlobalId for door in model.by_type("IfcDoor")}
    unscored = sorted(population - set(labels))
    counts = Counter(labels.values())
    return _case(
        "headroom_door_c321__mall", HEADROOM_DOOR_RULE, context.ifc_rel, "c.3.2.1", 1, labels, provenance,
        "reference_v2.py::build_headroom_doors (headroom.xlsx::Headroom_beta mall_UAT1_v6.4)",
        f"Doorway headroom of the {len(labels)} template doors: {dict(sorted(counts.items()))}; {len(unscored)} lift doors outside the template are unscored.",
        {"version": 1, "evaluation_scope": "annotated_subset", "unscored_ids": unscored, "population": len(population),
         "ifc_sha256": sha256(path),
         "status": "Advisor expert rows, each reproduced from the recorded ClearHeight under the authority's stated IFC convention.",
         "policy": "Doorway headroom is the recorded ClearHeight (authority convention); every template door is assessed.",
         "limitations": f"By OverallHeight, {sum(1 for p in provenance.values() if p['conflict'] and p['label'] == 'PASS')} PASS doors are below 2000 mm and {sum(1 for p in provenance.values() if p['conflict'] and p['label'] == 'FAIL')} FAIL doors are at least 2000 mm; the test model sets ClearHeight independently of the leaf size."},
    )


# --------------------------------------------------------------------------- #
# EC3-2026 catalogue: same-storey bottom elevation of spaces
# --------------------------------------------------------------------------- #
ELEVATION_TOLERANCE_MM = 10.0
ELEVATION_MODELS = ("mall", "scdf_3")
ELEVATION_RULE = (
    "Spaces in the same building storey must share the same bottom elevation.\n"
    "\nSource: BIM validation rule 'Same-storey elevation' of the EC3-2026 rule catalogue (general space "
    "check: spaces in the same storey share a bottom elevation).\n"
    "A space fails when the bottom of its modelled volume differs by more than 10 mm from the bottom "
    "elevation shared by most spaces of the storey it belongs to. Area and zone objects that are not rooms "
    "(IfcSpace with PredefinedType USERDEFINED, e.g. AREA_GFA or AREA_STRATA, or any IfcSpace carrying an "
    "IFC-SG area property set such as SGPset_SpaceArea_GFA) are not assessed: label them NA."
    + REPORT_LINE.format(entity="IfcSpace")
)


def build_space_elevation(model_id: str) -> GroundTruth:
    import ifcopenshell
    import ifcopenshell.util.element as element_util

    context = sr.MODELS[model_id]
    path = resolve_model_path(context.ifc_rel)
    model = ifcopenshell.open(str(path))
    spaces = model.by_type("IfcSpace")
    meshes = load_meshes(model, spaces)
    labels, provenance, storeys = {}, {}, {}
    for space in spaces:
        area_sets = sorted(name for name in element_util.get_psets(space) if name.startswith("SGPset_SpaceArea"))
        if str(space.PredefinedType) == "USERDEFINED" or area_sets:
            labels[space.GlobalId] = "NA"
            provenance[space.GlobalId] = {"space": _space_name(space), "object_type": space.ObjectType,
                                          "predefined_type": str(space.PredefinedType), "area_property_sets": area_sets,
                                          "label": "NA", "reason": "Area or zone object, not a room."}
            continue
        storey = element_util.get_aggregate(space)
        while storey is not None and not storey.is_a("IfcBuildingStorey"):
            storey = element_util.get_aggregate(storey) or element_util.get_container(storey)
        if storey is None or space.id() not in meshes:
            raise ValueError(f"{space.GlobalId}: no storey or geometry. Review required.")
        storeys.setdefault(storey, []).append((space, round(float(meshes[space.id()][0][:, 2].min()) * 1000.0, 1)))
    for storey, members in storeys.items():
        levels = Counter(round(z) for _, z in members)
        (shared, count), *rest = levels.most_common()
        if rest and rest[0][1] == count:
            raise ValueError(f"{storey.Name}: no single most common bottom elevation. Review required.")
        for space, z in members:
            difference = z - shared
            if 1.0 < abs(difference) <= 5 * ELEVATION_TOLERANCE_MM:
                raise ValueError(f"{space.GlobalId}: bottom differs by {difference:.1f} mm, too close to the tolerance. Review required.")
            label = "FAIL" if abs(difference) > ELEVATION_TOLERANCE_MM else "PASS"
            labels[space.GlobalId] = label
            provenance[space.GlobalId] = {"space": _space_name(space), "storey": storey.Name, "bottom_mm": z,
                                          "storey_shared_bottom_mm": shared, "spaces_in_storey": len(members),
                                          "spaces_at_shared_bottom": count, "difference_mm": round(difference, 1), "label": label}
    counts = Counter(labels.values())
    return _case(
        f"space_elevation__{model_id}", ELEVATION_RULE, context.ifc_rel, "BIM validation", 2, labels, provenance,
        "reference_v2.py::build_space_elevation",
        f"EC3 same-storey elevation rule on all IfcSpace: {dict(sorted(counts.items()))}.",
        {"version": 1, "evaluation_scope": "full_population", "population": len(labels), "ifc_sha256": sha256(path),
         "status": "Coded reference from IfcSpace geometry and spatial aggregation; no value lies between 1 mm and 50 mm from the shared level.",
         "policy": "Shared level = most common rounded bottom elevation of the storey's room spaces; tolerance 10 mm. USERDEFINED spaces and spaces carrying IFC-SG area property sets are NA.",
         "limitations": "A BIM consistency rule, not a building regulation: intentional raised or sunken floors (kitchens, wet areas) are reported as violations, as the rule states."},
    )


def expected_rule(case_id: str) -> str:
    """The builder's rule text for *case_id*, without regenerating the evidence."""
    if case_id.startswith("landing_interval_e352__"):
        return LANDING_RULE
    if case_id == "tread_width_e342__mall":
        return TREAD_RULE.format(building=sr.MODELS["mall"].building_phrase)
    if case_id.startswith("exit_requirements__"):
        return EXIT_RULE.format(requirement=template_requirement(EXIT_FOLDER, EXIT_MODELS[case_id.split("__")[1]] + ".xlsx"))
    return {"planter_box_gfa": PLANTER_RULE, "headroom_door_c321__mall": HEADROOM_DOOR_RULE}.get(case_id, ELEVATION_RULE)


def verify_recorded(case: GroundTruth, ifc_sha256: str) -> list[str]:
    """Fast consistency check of a stored usable_2 case; ``--check`` regenerates the evidence."""
    issues = []
    if case.case_id not in BUILDERS or case.reference.get("tier") != TIER:
        return ["Not a registered usable_2 case."]
    if case.reference.get("ifc_sha256") != ifc_sha256:
        issues.append("IFC fingerprint differs from the recorded reference.")
    if case.rule != expected_rule(case.case_id):
        issues.append("Rule text differs from the builder.")
    if set(case.provenance) != set(case.labels) or any(case.provenance[g].get("label") != label for g, label in case.labels.items()):
        issues.append("Recorded provenance and labels disagree.")
    return issues


# --------------------------------------------------------------------------- #
# Registry and command line
# --------------------------------------------------------------------------- #
BUILDERS = {
    **{f"landing_interval_e352__{m}": (lambda m=m: build_landing(m)) for m in LANDING_MODELS},
    "tread_width_e342__mall": build_tread_mall,
    **{f"exit_requirements__{m}": (lambda m=m: build_exit(m)) for m in EXIT_MODELS},
    "planter_box_gfa": build_planter,
    "headroom_door_c321__mall": build_headroom_doors,
    **{f"space_elevation__{m}": (lambda m=m: build_space_elevation(m)) for m in ELEVATION_MODELS},
}


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--only", action="append", choices=sorted(BUILDERS))
    parser.add_argument("--check", action="store_true", help="regenerate and compare with the stored cases")
    args = parser.parse_args()
    failed = False
    for case_id in args.only or sorted(BUILDERS):
        case = BUILDERS[case_id]()
        if args.check:
            same = load(case_id).as_dict() == json.loads(json.dumps(case.as_dict()))
            failed |= not same
            print(f"{case_id}: regenerated evidence {'matches' if same else 'DIFFERS FROM'} the stored case.")
        else:
            print(f"Saved {case_id}: {case.label_counts} to {case.save()}")
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
